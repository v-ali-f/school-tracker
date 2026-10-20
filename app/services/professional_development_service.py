"""Import, matching and three-year planning for staff courses."""

from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from hashlib import sha256
from io import BytesIO
import re

from openpyxl import load_workbook
from sqlalchemy import inspect, or_

from app.core.extensions import db
from app.models import StaffCourseImport, TeacherCourse, User
from app.utils.user_matching import normalize_fio


HEADERS = ("ФИО", "Должность", "Наименование курса", "Организация")
PEDAGOGICAL_ROLES = {
    "TEACHER", "CLASS_TEACHER", "DEPARTMENT_HEAD", "METHODIST",
    "SOCIAL_PEDAGOG", "PSYCHOLOGIST", "PRESCHOOL_TEACHER",
    "LOGOPEDIST", "DEFECTOLOGIST", "TUTOR",
}


def school_year_start(on_date=None):
    on_date = on_date or date.today()
    return on_date.year if on_date.month >= 9 else on_date.year - 1


def year_label(start):
    return f"{start}/{start + 1}"


def year_from_filename(filename):
    match = re.search(r"(?<!\d)(20\d{2})\s*[-–_/]\s*(20\d{2})(?!\d)", filename or "")
    if not match or int(match.group(2)) != int(match.group(1)) + 1:
        raise ValueError("Укажите учебный год в имени файла, например 2025-2026.xlsx")
    return int(match.group(1))


def parse_courses(content, filename):
    if not filename.lower().endswith(".xlsx"):
        raise ValueError("Нужен файл Excel .xlsx")
    if not content or len(content) > 16 * 1024 * 1024:
        raise ValueError("Файл пустой или больше 16 МБ")
    year = year_from_filename(filename)
    try:
        book = load_workbook(BytesIO(content), read_only=True, data_only=True)
        sheet = book.active
        iterator = sheet.iter_rows(values_only=True)
        headers = tuple(str(value or "").strip() for value in next(iterator)[:4])
        if headers != HEADERS:
            raise ValueError("Ожидаются столбцы: ФИО, Должность, Наименование курса, Организация")
        rows = []
        occurrences = Counter()
        for row_number, values in enumerate(iterator, start=2):
            if not any(value is not None and str(value).strip() for value in values):
                continue
            if len(values) < 4:
                raise ValueError(f"Строка {row_number}: недостаточно столбцов")
            name, position, title, provider = (str(value or "").strip() for value in values[:4])
            name = " ".join(name.split())
            title = " ".join(title.split())
            provider = " ".join(provider.split())
            if not name or not title:
                raise ValueError(f"Строка {row_number}: нужны ФИО и название курса")
            if any(len(value) > limit for value, limit in ((name, 255), (position, 255), (title, 512), (provider, 512))):
                raise ValueError(f"Строка {row_number}: слишком длинное значение")
            identity = (year, normalize_fio(fio=name), position.casefold(), title.casefold(), provider.casefold())
            occurrences[identity] += 1
            key = sha256(repr((*identity, occurrences[identity])).encode("utf-8")).hexdigest()
            rows.append({"source_key": key, "source_year": year, "source_row": row_number,
                         "employee_name": name, "position": position or None,
                         "source_title": title, "source_provider": provider or None})
        if not rows:
            raise ValueError("В таблице нет записей о курсах")
        return year, rows
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Не удалось прочитать файл Excel") from exc


def _user_lookup():
    by_name = defaultdict(list)
    for user in User.query.all():
        key = normalize_fio(fio=user.fio)
        if key:
            by_name[key].append(user)
    return by_name


def import_courses(content, filename, *, commit=True):
    """Replace one school's year snapshot, retaining reviewed user bindings."""
    year, rows = parse_courses(content, filename)
    source_name = filename.rsplit("/", 1)[-1].rsplit("\\", 1)[-1][:255]
    existing = {row.source_key: row for row in StaffCourseImport.query.filter_by(source_year=year).all()}
    users = _user_lookup()
    active_keys = set()
    created = linked = unlinked = 0
    for item in rows:
        active_keys.add(item["source_key"])
        row = existing.get(item["source_key"])
        if row is None:
            row = StaffCourseImport(source_key=item["source_key"])
            db.session.add(row)
            created += 1
        for key, value in item.items():
            setattr(row, key, value)
        row.source_name = source_name
        row.is_archived = False
        row.imported_at = datetime.utcnow()
        if row.teacher_id is None:
            candidates = users.get(normalize_fio(fio=row.employee_name), [])
            if len(candidates) == 1:
                row.teacher_id = candidates[0].id
        if row.teacher_id:
            linked += 1
        else:
            unlinked += 1
    archived = 0
    for key, row in existing.items():
        if key not in active_keys and not row.is_archived:
            row.is_archived = True
            archived += 1
    if commit:
        db.session.commit()
    return {"year": year, "total": len(rows), "created": created,
            "linked": linked, "unlinked": unlinked, "archived": archived}


def imported_and_manual_courses(user_ids=None):
    imported_query = StaffCourseImport.query.filter_by(is_archived=False)
    manual_query = TeacherCourse.query.filter_by(is_archived=False)
    if user_ids is not None:
        imported_query = imported_query.filter(StaffCourseImport.teacher_id.in_(user_ids))
        manual_query = manual_query.filter(TeacherCourse.teacher_id.in_(user_ids))
    result = defaultdict(list)
    for row in imported_query.all():
        if row.teacher_id:
            result[row.teacher_id].append(row)
    for row in manual_query.all():
        result[row.teacher_id].append(row)
    for rows in result.values():
        manual_keys = set()
        for row in rows:
            if isinstance(row, StaffCourseImport):
                continue
            when = row.end_date or row.start_date
            year = school_year_start(when) if when else (
                row.academic_year.start_date.year
                if row.academic_year and row.academic_year.start_date else None
            )
            if year is not None:
                manual_keys.add((year, row.title.casefold(), (row.provider or "").casefold()))
        rows[:] = [row for row in rows if not isinstance(row, StaffCourseImport)
                   or (row.source_year, row.title.casefold(), (row.provider or "").casefold())
                   not in manual_keys]
        rows.sort(key=_course_sort_key, reverse=True)
    return result


def _course_sort_key(row):
    exact = getattr(row, "completed_at", None) or row.end_date or row.start_date
    if exact:
        return (school_year_start(exact), exact.toordinal(), row.id)
    year = getattr(row, "source_year", None)
    if year is None and getattr(row, "academic_year", None):
        academic_year = row.academic_year
        year = academic_year.start_date.year if academic_year.start_date else None
    return (year or 0, 0, row.id)


def current_pedagogical_ids():
    """Use the same effective workload and membership rules as departments."""
    _, memberships = current_department_memberships()
    return set(memberships)


def current_department_memberships():
    """Map employees to departments in the current effective workload."""
    from app.models import AcademicYear, Department
    from app.services.workload_integration_service import department_teacher_ids

    departments = Department.query.order_by(Department.name).all()
    memberships = defaultdict(set)
    year = AcademicYear.query.filter_by(is_current=True).first()
    if year is None:
        return departments, memberships
    for department in departments:
        for teacher_id in department_teacher_ids(year.id, department.id):
            memberships[teacher_id].add(department.id)
    return departments, memberships


def current_preschool_teacher_ids():
    """Teachers assigned to active preschool groups, when that registry exists."""
    from app.models import AcademicYear
    from app.modules.preschool.models import PreschoolGroup

    inspector = inspect(db.engine)
    if not inspector.has_table("preschool_group"):
        return set()
    if "teacher_user_id" not in {
        column["name"] for column in inspector.get_columns("preschool_group")
    }:
        return set()
    year = AcademicYear.query.filter_by(is_current=True).first()
    query = PreschoolGroup.query.with_entities(PreschoolGroup.teacher_user_id).filter(
        PreschoolGroup.teacher_user_id.isnot(None),
        PreschoolGroup.is_active.is_(True),
        PreschoolGroup.is_archived.is_(False),
    )
    if year:
        query = query.filter(or_(
            PreschoolGroup.academic_year_id == year.id,
            PreschoolGroup.academic_year_id.is_(None),
        ))
    return {teacher_id for teacher_id, in query.all()}


def is_preschool_staff(user, courses=(), departments=(), group_teacher_ids=()):
    """Identify preschool staff from groups, role, imported position or department."""
    if user.id in group_teacher_ids:
        return True
    if "PRESCHOOL_TEACHER" in user.role_codes:
        return True
    if any(any(term in department.name.casefold() for term in ("дошколь", "детский сад"))
           for department in departments):
        return True
    return any(
        isinstance(row, StaffCourseImport)
        and any(term in (row.position or "").casefold()
                for term in ("воспитател", "дошколь", "детского сада"))
        for row in courses
    )


def is_pedagogical(user, courses=(), pedagogical_ids=None):
    if pedagogical_ids is not None and user.id in pedagogical_ids:
        return True
    if PEDAGOGICAL_ROLES.intersection(user.role_codes):
        return True
    return any(
        isinstance(row, StaffCourseImport)
        and any(word in (row.position or "").casefold()
                for word in ("учитель", "педагог", "воспитатель", "логопед", "психолог", "методист", "дефектолог", "тьютор", "инструктор", "музыкальный руководитель"))
        for row in courses
    )


def course_status(rows, today=None):
    """Planning status; course-year-only imports never imply an exact deadline."""
    today = today or date.today()
    if not rows:
        return {"code": "no_data", "label": "Нет сведений", "latest": None, "next": None}
    latest = max(rows, key=_course_sort_key)
    exact = getattr(latest, "completed_at", None) or latest.end_date or latest.start_date
    if exact:
        try:
            due_at = exact.replace(year=exact.year + 3)
        except ValueError:
            due_at = exact.replace(year=exact.year + 3, day=28)
        if due_at < today:
            code, label = "overdue", "Запланировать обучение"
        elif due_at <= today + timedelta(days=90):
            code, label = "due", "Срок подходит"
        else:
            code, label = "ok", "В пределах 3 лет"
        return {"code": code, "label": label, "latest": exact.strftime("%d.%m.%Y"),
                "next": due_at.strftime("%d.%m.%Y")}
    year = _course_sort_key(latest)[0]
    if not year:
        return {"code": "unknown", "label": "Уточнить дату", "latest": None, "next": None}
    due_year = year + 3
    current_year = school_year_start(today)
    if current_year > due_year:
        code, label = "overdue", "Запланировать обучение"
    elif current_year == due_year:
        code, label = "due", "Запланировать в этом году"
    else:
        code, label = "ok", "В пределах 3 лет"
    return {"code": code, "label": label, "latest": year_label(year),
            "next": year_label(due_year)}


def staff_overview(users, courses_by_user, today=None, pedagogical_ids=None):
    rows = []
    for user in users:
        courses = courses_by_user.get(user.id, [])
        status = course_status(courses, today=today)
        rows.append({"user": user, "courses": courses, "count": len(courses),
                     "pedagogical": is_pedagogical(user, courses, pedagogical_ids), "status": status})
    return rows
