"""Read administrator-provided МЭШ journal and ГИА workbook exports."""
import io
import re
import warnings
import zipfile
from collections import defaultdict
from datetime import date

from openpyxl import load_workbook


MONTHS = {"янв": 1, "фев": 2, "мар": 3, "апр": 4, "май": 5, "июн": 6,
          "июл": 7, "авг": 8, "сен": 9, "окт": 10, "ноя": 11, "дек": 12}
SUBJECTS = [
    ("русский язык", "Русский язык"), ("литератур", "Литература"),
    ("алгебр", "Алгебра"), ("геометр", "Геометрия"),
    ("вероятность", "Вероятность и статистика"), ("математик", "Математика"),
    ("физик", "Физика"), ("информатик", "Информатика"),
    ("биолог", "Биология"), ("хими", "Химия"),
    ("истори", "История"), ("обществознан", "Обществознание"),
    ("географ", "География"), ("ая ", "Английский язык"),
    ("английск", "Английский язык"), ("иностранный", "Английский язык"),
]


def _subject(title):
    lower = title.casefold()
    for part, name in SUBJECTS:
        if part in lower:
            return name
    return re.sub(r"\d{7,}$", "", title).strip()


def _normal(value):
    return " ".join(str(value).split())


def parse_journals_zip(content, filename):
    if len(content) > 64 * 1024 * 1024:
        raise ValueError("Архив превышает 64 МБ")
    records = {}
    stats = defaultdict(int)
    snapshots = set()
    warnings.simplefilter("ignore", UserWarning)
    try:
        archive = zipfile.ZipFile(io.BytesIO(content))
    except zipfile.BadZipFile as exc:
        raise ValueError("Выберите ZIP, скачанный из журналов МЭШ") from exc
    with archive:
        if sum(item.file_size for item in archive.infolist()) > 250 * 1024 * 1024:
            raise ValueError("Содержимое архива слишком велико")
        for name in archive.namelist():
            if not name.lower().endswith(".xlsx"):
                continue
            match = re.match(r"^(9|11)\s+\1[-\s]*([^\s]+).*?(\d{8})\.xlsx$", name, re.I)
            if not match:
                continue
            grade, letter, stamp = match.groups()
            cls = grade + letter.upper()
            snapshot = date.fromisoformat(f"{stamp[:4]}-{stamp[4:6]}-{stamp[6:]}")
            snapshots.add(snapshot)
            try:
                book = load_workbook(io.BytesIO(archive.read(name)), read_only=True, data_only=True)
            except Exception as exc:
                raise ValueError(f"Не удалось прочитать журнал {name}") from exc
            stats["files"] += 1
            for sheet in book:
                subject = _subject(sheet.title)
                rows = sheet.iter_rows(values_only=True)
                top = next(rows, ())
                days = next(rows, ())
                col_dates = {}
                month = None
                for col in range(2, min(len(top), len(days))):
                    header = str(top[col] or "").strip().casefold()
                    if header[:3] in MONTHS:
                        month = MONTHS[header[:3]]
                    try:
                        day = int(str(days[col]).strip())
                        if month is not None:
                            year = snapshot.year - (1 if snapshot.month < 9 and month >= 9 else 0)
                            value = date(year, month, day)
                            if value <= snapshot:
                                col_dates[col] = value.isoformat()
                    except (TypeError, ValueError):
                        pass
                if not col_dates:
                    continue
                stats["sheets"] += 1
                for row in rows:
                    if len(row) < 2 or not row[1] or not str(row[0] or "").strip().isdigit():
                        continue
                    student = _normal(row[1])
                    key = (cls, student, subject)
                    entry = records.setdefault(key, {"class": cls, "name": student,
                                                     "subject": subject, "marks": [], "absences": []})
                    for col, when in col_dates.items():
                        value = str(row[col] or "").strip() if col < len(row) else ""
                        if value in {"2", "3", "4", "5"}:
                            entry["marks"].append({"date": when, "grade": int(value)})
                            stats["marks"] += 1
                        elif value.casefold() in {"н", "нб"}:
                            entry["absences"].append(when)
                            stats["absences"] += 1
            book.close()
    if not records:
        raise ValueError("В архиве нет журналов 9 и 11 классов")
    if len(snapshots) != 1:
        raise ValueError("В архиве смешаны выгрузки с разными датами")
    return {"source": filename, "snapshot": snapshots.pop().isoformat(),
            "stats": dict(stats),
            "records": sorted(records.values(), key=lambda x: (x["class"], x["name"], x["subject"]))}


def parse_gia_xlsx(content, filename):
    if len(content) > 32 * 1024 * 1024:
        raise ValueError("Книга превышает 32 МБ")
    try:
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("Выберите XLSX со всеми листами таблицы ГИА") from exc
    students = []
    for sheet in book:
        cls = re.sub(r"\s+", "", sheet.title).upper()
        if not re.match(r"^(9|11)[А-Я]", cls):
            continue
        rows = sheet.iter_rows(values_only=True)
        next(rows, None)
        subjects = next(rows, ())
        kinds = next(rows, ())
        subject = ""
        columns = []
        for idx in range(2, max(len(subjects), len(kinds))):
            if idx < len(subjects) and subjects[idx]:
                subject = _normal(subjects[idx])
            kind = str(kinds[idx] if idx < len(kinds) and kinds[idx] is not None else "").strip()
            if re.fullmatch(r"(9|11)\s*класс", kind, re.I):
                columns.append((idx, subject))
        for row in rows:
            name = _normal(row[1]) if len(row) > 1 and row[1] else ""
            number = row[0] if row else None
            valid_number = (number >= 1 and number == int(number)) if isinstance(number, (int, float)) and not isinstance(number, bool) else str(number or "").strip().isdigit()
            if len(name.split()) < 2 or not valid_number:
                continue
            choices = []
            for idx, subject in columns:
                cell = row[idx] if idx < len(row) else None
                value = str(cell).strip() if cell is not None else ""
                selected = (cell == 1 and not isinstance(cell, bool)) or value == "1"
                if selected or value.lower() in {"база", "проф", "профиль"}:
                    item = {"subject": subject}
                    if not selected:
                        item["level"] = value
                    choices.append(item)
            students.append({"class": cls, "name": name, "choices": choices})
    book.close()
    if not students:
        raise ValueError("В книге не найдены ученики 9 или 11 классов")
    return {"source": filename, "asOf": date.today().isoformat(), "students": students}
