import os
from datetime import date, datetime, time as dt_time, timedelta
from io import BytesIO
import calendar as pycalendar
from xml.sax.saxutils import escape

from flask import Blueprint, abort, flash, redirect, render_template, request, url_for, send_file
from flask_login import current_user, login_required
from sqlalchemy import and_, func, or_

from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

from app.core.extensions import db
from app.models import (
    AcademicYear,
    Building,
    Department,
    SchoolClass,
    SchoolPlanCategory,
    SchoolPlanDirection,
    SchoolPlanEditorAccess,
    SchoolPlanEvent,
    SchoolPlanEventClass,
    SchoolPlanEventGrade,
    SchoolPlanEventGroup,
    SchoolPlanEventResponsible,
    User,
)
from app.services.school_plan_access import (
    can_fill_school_plan,
    can_manage_school_plan_editors,
    has_implicit_plan_access,
    has_protected_plan_access,
)
from app.services.school_plan_service import ensure_school_plan_seed_data

school_plan_bp = Blueprint('school_plan', __name__, url_prefix='/school-plan')


STATUS_TITLES = {
    'planned': 'Запланировано',
    'in_progress': 'В работе',
    'done': 'Выполнено',
    'postponed': 'Перенесено',
    'cancelled': 'Отменено',
}

PRIORITY_TITLES = {
    'normal': 'Обычный',
    'important': 'Важный',
    'critical': 'Очень важный',
}

VISIBILITY_TITLES = {
    'school': 'Вся школа',
    'building': 'Здание',
    'class': 'Класс',
}

WEEKDAY_TITLES = ['Понедельник', 'Вторник', 'Среда', 'Четверг', 'Пятница', 'Суббота', 'Воскресенье']
MONTH_TITLES = [
    '', 'Январь', 'Февраль', 'Март', 'Апрель', 'Май', 'Июнь',
    'Июль', 'Август', 'Сентябрь', 'Октябрь', 'Ноябрь', 'Декабрь'
]
MONTH_TITLES_GENITIVE = [
    '', 'января', 'февраля', 'марта', 'апреля', 'мая', 'июня',
    'июля', 'августа', 'сентября', 'октября', 'ноября', 'декабря'
]
WEEKDAY_SHORT_TITLES = ['ПН', 'ВТ', 'СР', 'ЧТ', 'ПТ', 'СБ', 'ВС']

SCHOOL_PLAN_POSITION_OPTIONS = [
    ('TEACHER', 'Педагоги'),
    ('METHODIST', 'Методисты'),
    ('DEPUTY_DIRECTOR', 'Заместители директора'),
    ('PEDAGOG_ORGANIZER', 'Педагоги-организаторы'),
    ('SOCIAL_PEDAGOG', 'Социальные педагоги'),
    ('EDUCATOR', 'Воспитатели'),
    ('SENIOR_EDUCATOR', 'Старшие воспитатели'),
    ('PSYCHOLOGIST', 'Педагоги-психологи'),
    ('LOGOPEDIST', 'Учителя-логопеды'),
    ('DEFECTOLOGIST', 'Учителя-дефектологи'),
    ('TUTOR', 'Тьюторы'),
    ('ASSISTANT', 'Ассистенты'),
]

SCHOOL_PLAN_ROLE_LABELS = {
    'ADMIN': 'Администратор',
    'DIRECTOR': 'Директор',
    'DEPUTY_DIRECTOR': 'Заместитель директора',
    'METHODIST': 'Методист',
    'TEACHER': 'Педагог',
    'CLASS_TEACHER': 'Классный руководитель',
    'PEDAGOG_ORGANIZER': 'Педагог-организатор',
    'SOCIAL_PEDAGOG': 'Социальный педагог',
    'EDUCATOR': 'Воспитатель',
    'SENIOR_EDUCATOR': 'Старший воспитатель',
    'PSYCHOLOGIST': 'Педагог-психолог',
    'LOGOPEDIST': 'Учитель-логопед',
    'DEFECTOLOGIST': 'Учитель-дефектолог',
    'TUTOR': 'Тьютор',
    'ASSISTANT': 'Ассистент',
}

BUILDING_COLORS = [
    '#2563eb', '#059669', '#dc2626', '#7c3aed', '#d97706',
    '#0891b2', '#be185d', '#4f46e5', '#65a30d', '#9333ea',
]


PDF_FONT_NAME = None
PDF_FONT_BOLD_NAME = None


class SchoolPlanFormError(ValueError):
    pass


def _find_pdf_font_paths():
    regular_candidates = [
        '/System/Library/Fonts/Supplemental/Arial.ttf',
        '/Library/Fonts/Arial.ttf',
        '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',
        '/usr/share/fonts/dejavu/DejaVuSans.ttf',
        'C:/Windows/Fonts/arial.ttf',
        'C:/Windows/Fonts/DejaVuSans.ttf',
    ]
    bold_candidates = [
        '/System/Library/Fonts/Supplemental/Arial Bold.ttf',
        '/Library/Fonts/Arial Bold.ttf',
        '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf',
        '/usr/share/fonts/dejavu/DejaVuSans-Bold.ttf',
        'C:/Windows/Fonts/arialbd.ttf',
        'C:/Windows/Fonts/DejaVuSans-Bold.ttf',
    ]

    regular_path = next((path for path in regular_candidates if os.path.exists(path)), None)
    bold_path = next((path for path in bold_candidates if os.path.exists(path)), None)
    return regular_path, bold_path



def _ensure_pdf_fonts():
    global PDF_FONT_NAME, PDF_FONT_BOLD_NAME
    if PDF_FONT_NAME and PDF_FONT_BOLD_NAME:
        return PDF_FONT_NAME, PDF_FONT_BOLD_NAME

    regular_path, bold_path = _find_pdf_font_paths()
    if regular_path:
        if 'SchoolPlanPdfFont' not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont('SchoolPlanPdfFont', regular_path))
        PDF_FONT_NAME = 'SchoolPlanPdfFont'
    else:
        PDF_FONT_NAME = 'Helvetica'

    if bold_path:
        if 'SchoolPlanPdfFontBold' not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont('SchoolPlanPdfFontBold', bold_path))
        PDF_FONT_BOLD_NAME = 'SchoolPlanPdfFontBold'
    else:
        PDF_FONT_BOLD_NAME = PDF_FONT_NAME

    return PDF_FONT_NAME, PDF_FONT_BOLD_NAME


def _is_manager():
    return can_fill_school_plan(current_user)


def _can_manage_editors():
    return can_manage_school_plan_editors(current_user)


def _user_scope():
    if _is_manager():
        return {'all_access': True, 'building_ids': set(), 'class_ids': set(), 'grades': set()}

    building_ids = {row.building_id for row in getattr(current_user, 'building_links', []) if getattr(row, 'building_id', None)}
    class_rows = SchoolClass.query.filter_by(teacher_user_id=current_user.id, is_archived=False).all()
    class_ids = {row.id for row in class_rows}
    grades = {row.grade for row in class_rows if row.grade is not None}
    for row in class_rows:
        if row.building_id:
            building_ids.add(row.building_id)

    return {
        'all_access': False,
        'building_ids': building_ids,
        'class_ids': class_ids,
        'grades': grades,
    }


def _visible_events_query(include_archived=False):
    query = SchoolPlanEvent.query
    if not include_archived:
        query = query.filter(SchoolPlanEvent.is_archived.is_(False))

    scope = _user_scope()
    if scope['all_access']:
        return query

    conditions = [SchoolPlanEvent.visibility_level == 'school']
    if scope['building_ids']:
        conditions.append(and_(SchoolPlanEvent.visibility_level == 'building', SchoolPlanEvent.building_id.in_(scope['building_ids'])))
    if scope['class_ids']:
        conditions.append(and_(SchoolPlanEvent.visibility_level == 'class', SchoolPlanEvent.class_id.in_(scope['class_ids'])))
        conditions.append(
            SchoolPlanEvent.target_class_links.any(
                SchoolPlanEventClass.class_id.in_(scope['class_ids'])
            )
        )
    if scope['grades']:
        conditions.append(
            SchoolPlanEvent.target_grade_links.any(
                SchoolPlanEventGrade.grade.in_(scope['grades'])
            )
        )

    return query.filter(or_(*conditions))


def _parse_date(value):
    if not value:
        return None
    try:
        return datetime.strptime(value, '%Y-%m-%d').date()
    except Exception:
        return None


def _apply_filters(query):
    year_id = _args_int('year_id')
    building_id = _args_int('building_id')
    class_id = _args_int('class_id')
    direction_id = _args_int('direction_id')
    category_id = _args_int('category_id')
    visibility_level = (request.args.get('visibility_level') or '').strip()
    status = (request.args.get('status') or '').strip()
    priority = (request.args.get('priority') or '').strip()
    q = (request.args.get('q') or '').strip()

    if year_id:
        query = query.filter(SchoolPlanEvent.academic_year_id == year_id)
    if building_id:
        query = query.filter(SchoolPlanEvent.building_id == building_id)
    if class_id:
        query = query.filter(SchoolPlanEvent.class_id == class_id)
    if direction_id:
        query = query.filter(SchoolPlanEvent.direction_id == direction_id)
    if category_id:
        query = query.filter(SchoolPlanEvent.category_id == category_id)
    if visibility_level:
        query = query.filter(SchoolPlanEvent.visibility_level == visibility_level)
    if status:
        query = query.filter(SchoolPlanEvent.status == status)
    if priority:
        query = query.filter(SchoolPlanEvent.priority == priority)
    if q:
        ilike = f'%{q}%'
        query = query.filter(or_(
            SchoolPlanEvent.title.ilike(ilike),
            SchoolPlanEvent.description.ilike(ilike),
            SchoolPlanEvent.responsible_text.ilike(ilike),
            SchoolPlanEvent.participants.ilike(ilike),
            SchoolPlanEvent.responsible_links.any(
                SchoolPlanEventResponsible.user.has(
                    or_(
                        User.last_name.ilike(ilike),
                        User.first_name.ilike(ilike),
                        User.middle_name.ilike(ilike),
                        User.username.ilike(ilike),
                    )
                )
            ),
        ))
    return query


def _event_query(include_archived=False):
    return _apply_filters(_visible_events_query(include_archived=include_archived))


def _period_overlap(query, start_date, end_date):
    return query.filter(
        SchoolPlanEvent.start_date <= end_date,
        func.coalesce(SchoolPlanEvent.end_date, SchoolPlanEvent.start_date) >= start_date,
    )


def _selected_month():
    today = date.today()
    raw_period = (request.args.get('period') or '').strip()
    if raw_period:
        try:
            parsed = datetime.strptime(raw_period, '%Y-%m')
            return parsed.year, parsed.month
        except ValueError:
            pass
    year = _args_int('year') or today.year
    month = _args_int('month') or today.month
    return _month_shift(year, month, 0)


def _normalize_period(value):
    raw = (value or '').strip()
    try:
        parsed = datetime.strptime(raw, '%Y-%m')
    except ValueError:
        return None
    return f'{parsed.year:04d}-{parsed.month:02d}'


def _parse_time(value):
    raw = (value or '').strip()
    if not raw:
        return None
    try:
        return datetime.strptime(raw, '%H:%M').time()
    except ValueError as exc:
        raise SchoolPlanFormError('Укажите корректное время начала.') from exc


def _event_ordering():
    return (
        SchoolPlanEvent.start_date.asc(),
        SchoolPlanEvent.start_time.asc().nullslast(),
        SchoolPlanEvent.created_at.desc(),
    )


def _event_sort_key(event):
    return (
        event.start_date,
        event.start_time is None,
        event.start_time or dt_time.max,
        (event.title or '').lower(),
    )


def _month_filter_options(year, month):
    options = []
    for option_year in range(year - 1, year + 2):
        for option_month in range(1, 13):
            options.append({
                'value': f'{option_year:04d}-{option_month:02d}',
                'label': f'{MONTH_TITLES[option_month]} {option_year}',
            })
    return options


def _return_period(event=None):
    selected = _normalize_period(
        request.form.get('return_period') or request.args.get('return_period')
    )
    if selected:
        return selected
    if event and event.start_date:
        return event.start_date.strftime('%Y-%m')
    today = date.today()
    return f'{today.year:04d}-{today.month:02d}'


def _month_bounds(year, month):
    start = date(year, month, 1)
    end = date(year, month, pycalendar.monthrange(year, month)[1])
    return start, end


def _month_query(query, year, month):
    start, end = _month_bounds(year, month)
    return _period_overlap(query, start, end)


def _month_shift(year, month, delta):
    month += delta
    while month < 1:
        month += 12
        year -= 1
    while month > 12:
        month -= 12
        year += 1
    return year, month


def _build_calendar_weeks(events, year, month):
    cal = pycalendar.Calendar(firstweekday=0)
    event_map = {}
    for event in events:
        cur = max(event.start_date, date(year, month, 1))
        last_day = pycalendar.monthrange(year, month)[1]
        end = min(event.end_date or event.start_date, date(year, month, last_day))
        while cur <= end:
            event_map.setdefault(cur, []).append(event)
            cur += timedelta(days=1)

    weeks = []
    for week in cal.monthdatescalendar(year, month):
        days = []
        for day in week:
            days.append({
                'date': day,
                'in_month': day.month == month,
                'events': sorted(event_map.get(day, []), key=_event_sort_key),
                'is_today': day == date.today(),
            })
        weeks.append(days)
    return weeks


def _agenda_week_groups(events, year, month):
    month_start, month_end = _month_bounds(year, month)
    grouped = {}
    for event in events:
        anchor = max(event.start_date, month_start)
        week_start = anchor - timedelta(days=anchor.weekday())
        visible_start = max(week_start, month_start)
        visible_end = min(week_start + timedelta(days=6), month_end)
        grouped.setdefault((visible_start, visible_end), []).append(event)

    result = []
    for (visible_start, visible_end), rows in sorted(grouped.items()):
        if visible_start == visible_end:
            label = f'{visible_start.day} {MONTH_TITLES_GENITIVE[visible_start.month]}'
        elif visible_start.month == visible_end.month:
            label = (
                f'{visible_start.day}–{visible_end.day} '
                f'{MONTH_TITLES_GENITIVE[visible_end.month]}'
            )
        else:
            label = (
                f'{visible_start.day} {MONTH_TITLES_GENITIVE[visible_start.month]} — '
                f'{visible_end.day} {MONTH_TITLES_GENITIVE[visible_end.month]}'
            )
        result.append({
            'start': visible_start,
            'end': visible_end,
            'label': label,
            'events': sorted(rows, key=_event_sort_key),
        })
    return result


def _school_plan_sidebar_context():
    from app.modules.hub.routes import build_home_context

    home_context = build_home_context()
    page = home_context.get('page') or {}
    seen_urls = set()

    def unique_items(items, limit=None):
        result = []
        for item in items or []:
            url = item.get('url')
            if (
                not url
                or url in seen_urls
                or item.get('endpoint') == 'school_plan.index'
            ):
                continue
            seen_urls.add(url)
            result.append(item)
            if limit and len(result) >= limit:
                break
        return result

    return {
        'workspace_nav': home_context,
        'workspace_context_title': 'Рабочее пространство',
        'workspace_context_subtitle': 'Планирование мероприятий школы',
        'plan_sidebar_daily_actions': unique_items(page.get('quick_actions')),
        'plan_sidebar_sections': unique_items(page.get('secondary_sections')),
        'plan_sidebar_service_actions': unique_items(page.get('admin_sections')),
    }



def _args_int(name):
    value = (request.args.get(name) or '').strip()
    try:
        return int(value) if value else None
    except Exception:
        return None


def _request_args_dict():
    return request.args.to_dict(flat=True)


def _export_filename(prefix, ext):
    stamp = datetime.now().strftime('%Y%m%d_%H%M')
    return f'{prefix}_{stamp}.{ext}'


def _events_for_export(include_archived=False):
    year, month = _selected_month()
    query = _month_query(_event_query(include_archived=include_archived), year, month)
    return query.order_by(*_event_ordering()).all()


def _build_excel(events):
    wb = Workbook()
    ws = wb.active
    ws.title = 'План работы школы'
    headers = ['Дата / время', 'Мероприятие', 'Направление', 'Классы', 'Ответственные', 'Примечание']
    ws.append(headers)
    for cell in ws[1]:
        cell.font = cell.font.copy(bold=True)
    for event in events:
        ws.append([
            event.display_schedule,
            event.title or '',
            event.direction.name if event.direction else '',
            event.display_audience,
            event.display_responsible or '',
            event.description or '',
        ])
    widths = [24, 52, 28, 32, 36, 45]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64+i)].width = width
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = cell.alignment.copy(wrap_text=True, vertical='top')
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream


def _build_pdf(events, title='План работы школы'):
    font_name, bold_font_name = _ensure_pdf_fonts()

    stream = BytesIO()
    doc = SimpleDocTemplate(stream, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=24, bottomMargin=24)
    styles = getSampleStyleSheet()
    styles['Title'].fontName = bold_font_name
    styles['Title'].fontSize = 19
    styles['Title'].leading = 23
    cell_style = ParagraphStyle(
        'SchoolPlanPdfCell',
        parent=styles['BodyText'],
        fontName=font_name,
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor('#1f2937'),
        spaceBefore=0,
        spaceAfter=0,
        splitLongWords=True,
    )
    header_style = ParagraphStyle(
        'SchoolPlanPdfHeader',
        parent=cell_style,
        fontName=bold_font_name,
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor('#111827'),
    )

    def cell(value, style=cell_style):
        text = escape(str(value or '')).replace('\n', '<br/>')
        return Paragraph(text or '&nbsp;', style)

    story = [Paragraph(title, styles['Title']), Spacer(1, 10)]
    data = [[
        cell('Дата / время', header_style),
        cell('Мероприятие', header_style),
        cell('Направление', header_style),
        cell('Классы', header_style),
        cell('Ответственные', header_style),
    ]]
    for event in events:
        data.append([
            cell(event.display_schedule),
            cell(event.title),
            cell(event.direction.name if event.direction else ''),
            cell(event.display_audience),
            cell(event.display_responsible),
        ])
    table = Table(
        data,
        repeatRows=1,
        colWidths=[88, 225, 120, 145, 215],
        hAlign='LEFT',
        splitByRow=1,
    )
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#e2e8f0')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.black),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
        ('FONTNAME', (0, 0), (-1, 0), bold_font_name),
        ('FONTNAME', (0, 1), (-1, -1), font_name),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
        ('LEFTPADDING', (0, 0), (-1, -1), 6),
        ('RIGHTPADDING', (0, 0), (-1, -1), 6),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ]))
    story.append(table)
    doc.build(story)
    stream.seek(0)
    return stream


def _academic_week_groups(events):
    groups = {}
    for event in events:
        week_no = event.start_date.isocalendar().week
        groups.setdefault(week_no, []).append(event)
    ordered = []
    for week_no in sorted(groups):
        ordered.append((week_no, sorted(groups[week_no], key=_event_sort_key)))
    return ordered


@school_plan_bp.route('/')
@login_required
def index():
    include_archived = _is_manager() and request.args.get('show_archived') == '1'
    year, month = _selected_month()
    events = (
        _month_query(_event_query(include_archived=include_archived), year, month)
        .order_by(*_event_ordering())
        .all()
    )
    prev_year, prev_month = _month_shift(year, month, -1)
    next_year, next_month = _month_shift(year, month, 1)
    month_start, month_end = _month_bounds(year, month)
    all_month_events = (
        _month_query(_event_query(include_archived=True), year, month)
        .order_by(*_event_ordering())
        .all()
    )
    today = date.today()
    current_week_start = today - timedelta(days=today.weekday())
    current_week_end = current_week_start + timedelta(days=6)
    long_events = [
        event for event in events
        if event.end_date and event.end_date > event.start_date
    ]
    direction_counts = {}
    for event in events:
        if not event.direction:
            continue
        bucket = direction_counts.setdefault(event.direction_id, {
            'direction': event.direction,
            'count': 0,
        })
        bucket['count'] += 1
    upcoming_event = next(
        (
            event for event in events
            if (event.end_date or event.start_date) >= today
        ),
        events[0] if events else None,
    )
    return render_template(
        'school_plan/index.html',
        events=events,
        can_manage=_is_manager(),
        active_view='list',
        year=year,
        month=month,
        selected_period=f'{year:04d}-{month:02d}',
        month_title=MONTH_TITLES[month],
        prev_period=f'{prev_year:04d}-{prev_month:02d}',
        next_period=f'{next_year:04d}-{next_month:02d}',
        agenda_groups=_agenda_week_groups(events, year, month),
        calendar_weeks=_build_calendar_weeks(events, year, month),
        plan_summary={
            'total': len(events),
            'this_week': sum(
                1 for event in events
                if event.start_date <= current_week_end
                and (event.end_date or event.start_date) >= current_week_start
            ),
            'long': len(long_events),
            'cancelled': sum(
                1 for event in all_month_events
                if event.is_archived or event.status == 'cancelled'
            ),
        },
        direction_legend=sorted(
            direction_counts.values(),
            key=lambda row: (-row['count'], row['direction'].name.lower()),
        ),
        upcoming_event=upcoming_event,
        month_start=month_start,
        month_end=month_end,
        month_titles_genitive=MONTH_TITLES_GENITIVE,
        weekday_short_titles=WEEKDAY_SHORT_TITLES,
        role_labels=SCHOOL_PLAN_ROLE_LABELS,
        **_school_plan_sidebar_context(),
        **_form_context(),
    )


@school_plan_bp.route('/week')
@login_required
def week_view():
    ref = _parse_date(request.args.get('date')) or date.today()
    week_start = ref - timedelta(days=ref.weekday())
    week_end = week_start + timedelta(days=6)
    include_archived = _is_manager() and request.args.get('show_archived') == '1'
    events = _period_overlap(_event_query(include_archived=include_archived), week_start, week_end).order_by(*_event_ordering()).all()
    days = []
    for i in range(7):
        day = week_start + timedelta(days=i)
        day_events = [e for e in events if e.start_date <= day <= (e.end_date or e.start_date)]
        days.append({'date': day, 'title': WEEKDAY_TITLES[i], 'events': day_events})
    return render_template('school_plan/week.html', days=days, week_start=week_start, week_end=week_end,
                           prev_date=week_start - timedelta(days=7), next_date=week_start + timedelta(days=7),
                           can_manage=_is_manager(), active_view='week', **_form_context())


@school_plan_bp.route('/month')
@login_required
def month_view():
    year, month = _selected_month()
    month_start, month_end = _month_bounds(year, month)
    include_archived = _is_manager() and request.args.get('show_archived') == '1'
    events = _period_overlap(_event_query(include_archived=include_archived), month_start, month_end).order_by(*_event_ordering()).all()
    weeks = _build_calendar_weeks(events, year, month)
    prev_y, prev_m = _month_shift(year, month, -1)
    next_y, next_m = _month_shift(year, month, 1)
    return render_template('school_plan/month.html', weeks=weeks, month=month, year=year,
                           month_title=MONTH_TITLES[month], prev_year=prev_y, prev_month=prev_m,
                           next_year=next_y, next_month=next_m,
                           selected_period=f'{year:04d}-{month:02d}',
                           can_manage=_is_manager(), active_view='month', **_form_context())


@school_plan_bp.route('/day')
@login_required
def day_view():
    selected_date = _parse_date(request.args.get('date')) or date.today()
    include_archived = _is_manager() and request.args.get('show_archived') == '1'
    events = _period_overlap(_event_query(include_archived=include_archived), selected_date, selected_date).order_by(*_event_ordering()).all()
    return render_template('school_plan/day.html', selected_date=selected_date, events=events,
                           prev_date=selected_date - timedelta(days=1), next_date=selected_date + timedelta(days=1),
                           can_manage=_is_manager(), active_view='day', **_form_context())


@school_plan_bp.route('/weeks')
@login_required
def weeks_view():
    include_archived = _is_manager() and request.args.get('show_archived') == '1'
    events = _event_query(include_archived=include_archived).order_by(*_event_ordering()).all()
    groups = _academic_week_groups(events)
    return render_template('school_plan/weeks.html', week_groups=groups, can_manage=_is_manager(), active_view='weeks', **_form_context())


@school_plan_bp.route('/new', methods=['GET', 'POST'])
@login_required
def create():
    if not _is_manager():
        abort(403)
    event = SchoolPlanEvent(start_date=_parse_date(request.args.get('date')))
    return_period = _normalize_period(
        request.form.get('return_period') or request.args.get('return_period')
    )
    if request.method == 'POST':
        try:
            _apply_event_form(event)
            db.session.add(event)
            db.session.commit()
        except SchoolPlanFormError as exc:
            db.session.rollback()
            flash(str(exc), 'danger')
        else:
            flash('Мероприятие плана создано.', 'success')
            return redirect(
                url_for(
                    'school_plan.index',
                    period=return_period or event.start_date.strftime('%Y-%m'),
                )
            )
    return render_template(
        'school_plan/form.html', event=event, return_period=return_period or '',
        can_manage=True, active_view='list',
        **_school_plan_sidebar_context(), **_form_context(event)
    )


@school_plan_bp.route('/<int:event_id>')
@login_required
def view(event_id):
    event = _visible_events_query(include_archived=True).filter(SchoolPlanEvent.id == event_id).first_or_404()
    return render_template('school_plan/view.html', event=event, can_manage=_is_manager(), active_view='list',
                           status_titles=STATUS_TITLES, priority_titles=PRIORITY_TITLES,
                           visibility_titles=VISIBILITY_TITLES, **_school_plan_sidebar_context())


@school_plan_bp.route('/<int:event_id>/edit', methods=['GET', 'POST'])
@login_required
def edit(event_id):
    if not _is_manager():
        abort(403)
    event = SchoolPlanEvent.query.get_or_404(event_id)
    return_period = _normalize_period(
        request.form.get('return_period') or request.args.get('return_period')
    ) or event.start_date.strftime('%Y-%m')
    if request.method == 'POST':
        try:
            _apply_event_form(event)
            db.session.commit()
        except SchoolPlanFormError as exc:
            db.session.rollback()
            flash(str(exc), 'danger')
        else:
            flash('Мероприятие плана обновлено.', 'success')
            return redirect(url_for('school_plan.index', period=return_period))
    return render_template(
        'school_plan/form.html', event=event, return_period=return_period,
        can_manage=True, active_view='list',
        **_school_plan_sidebar_context(), **_form_context(event)
    )


@school_plan_bp.route('/<int:event_id>/archive', methods=['POST'])
@login_required
def archive(event_id):
    if not _is_manager():
        abort(403)
    event = SchoolPlanEvent.query.get_or_404(event_id)
    event.is_archived = True
    event.updated_by_user_id = current_user.id
    db.session.commit()
    flash('Мероприятие отменено и перенесено в архив.', 'success')
    return redirect(url_for('school_plan.index', period=_return_period(event)))


@school_plan_bp.route('/<int:event_id>/restore', methods=['POST'])
@login_required
def restore(event_id):
    if not _is_manager():
        abort(403)
    event = SchoolPlanEvent.query.get_or_404(event_id)
    event.is_archived = False
    event.updated_by_user_id = current_user.id
    db.session.commit()
    flash('Мероприятие восстановлено из архива.', 'success')
    return redirect(
        url_for(
            'school_plan.index', period=_return_period(event), show_archived=1
        )
    )


@school_plan_bp.route('/<int:event_id>/delete', methods=['POST'])
@login_required
def delete(event_id):
    if not _is_manager():
        abort(403)
    event = SchoolPlanEvent.query.get_or_404(event_id)
    return_period = _return_period(event)
    db.session.delete(event)
    db.session.commit()
    flash('Мероприятие плана удалено.', 'success')
    return redirect(url_for('school_plan.index', period=return_period))



@school_plan_bp.route('/export/xlsx')
@login_required
def export_xlsx():
    include_archived = _is_manager() and request.args.get('show_archived') == '1'
    year, month = _selected_month()
    stream = _build_excel(_events_for_export(include_archived=include_archived))
    return send_file(stream,
                     as_attachment=True,
                     download_name=f'school_plan_{year:04d}_{month:02d}.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@school_plan_bp.route('/export/pdf')
@login_required
def export_pdf():
    include_archived = _is_manager() and request.args.get('show_archived') == '1'
    year, month = _selected_month()
    title = f'План работы школы - {MONTH_TITLES[month].lower()} {year}'
    stream = _build_pdf(_events_for_export(include_archived=include_archived), title=title)
    return send_file(stream,
                     as_attachment=True,
                     download_name=f'school_plan_{year:04d}_{month:02d}.pdf',
                     mimetype='application/pdf')


@school_plan_bp.route('/print')
@login_required
def print_view_export():
    include_archived = _is_manager() and request.args.get('show_archived') == '1'
    year, month = _selected_month()
    events = _events_for_export(include_archived=include_archived)
    return render_template('school_plan/print.html', events=events,
                           export_title=f'План работы школы — {MONTH_TITLES[month].lower()} {year}',
                           status_titles=STATUS_TITLES, priority_titles=PRIORITY_TITLES,
                           visibility_titles=VISIBILITY_TITLES)


@school_plan_bp.route('/editors', methods=['GET', 'POST'])
@login_required
def editors():
    if not _can_manage_editors():
        abort(403)

    if request.method == 'POST':
        action = (request.form.get('action') or '').strip()
        if action == 'add':
            user_id = request.form.get('user_id', type=int)
            user = User.query.filter_by(id=user_id, is_active_user=True).first()
            if not user:
                flash('Сотрудник не найден или его учётная запись неактивна.', 'danger')
            elif has_protected_plan_access(user):
                flash('Доступ директора к управлению планом является защищённым.', 'info')
            else:
                access = SchoolPlanEditorAccess.query.filter_by(user_id=user.id).first()
                if access and access.is_enabled:
                    flash('Сотрудник уже имеет доступ к плану.', 'info')
                    return redirect(url_for('school_plan.editors'))
                if access:
                    access.is_enabled = True
                    access.granted_by_user_id = current_user.id
                else:
                    db.session.add(
                        SchoolPlanEditorAccess(
                            user_id=user.id,
                            granted_by_user_id=current_user.id,
                            is_enabled=True,
                        )
                    )
                db.session.commit()
                flash(f'Сотруднику «{user.fio or user.username}» предоставлен доступ.', 'success')
        elif action == 'remove':
            user_id = request.form.get('user_id', type=int)
            user = User.query.filter_by(id=user_id, is_active_user=True).first_or_404()
            if has_protected_plan_access(user):
                flash('Нельзя снять защищённый доступ директора к управлению планом.', 'danger')
                return redirect(url_for('school_plan.editors'))
            access = SchoolPlanEditorAccess.query.filter_by(user_id=user.id).first()
            if access:
                access.is_enabled = False
                access.granted_by_user_id = current_user.id
            else:
                db.session.add(
                    SchoolPlanEditorAccess(
                        user_id=user.id,
                        granted_by_user_id=current_user.id,
                        is_enabled=False,
                    )
                )
            db.session.commit()
            flash(f'Доступ сотрудника «{user.fio or user.username}» снят.', 'success')
        else:
            flash('Не удалось определить действие.', 'danger')
        return redirect(url_for('school_plan.editors'))

    all_active_users = User.query.filter(User.is_active_user.is_(True)).order_by(
        User.last_name.asc(), User.first_name.asc(), User.middle_name.asc()
    ).all()
    access_by_user_id = {
        row.user_id: row for row in SchoolPlanEditorAccess.query.all()
    }
    assignments = []
    available_users = []
    for user in all_active_users:
        access = access_by_user_id.get(user.id)
        protected = has_protected_plan_access(user)
        implicit = has_implicit_plan_access(user)
        effective = protected or (bool(access.is_enabled) if access else implicit)
        if not effective:
            available_users.append(user)
            continue
        assignments.append(
            {
                'user': user,
                'access': access,
                'source': (
                    'protected' if protected
                    else 'role' if implicit
                    else 'assigned'
                ),
                'can_remove': not protected,
            }
        )
    editor_summary = {
        'total': len(assignments),
        'assigned': sum(1 for entry in assignments if entry['source'] == 'assigned'),
        'role': sum(1 for entry in assignments if entry['source'] == 'role'),
        'protected': sum(1 for entry in assignments if entry['source'] == 'protected'),
        'available': len(available_users),
    }
    return render_template(
        'school_plan/editors.html',
        assignments=assignments,
        available_users=available_users,
        editor_summary=editor_summary,
        role_labels=SCHOOL_PLAN_ROLE_LABELS,
        can_manage=True,
        can_manage_editors=True,
        active_view='editors',
        current_args={},
        **_school_plan_sidebar_context(),
    )


@school_plan_bp.route('/legend')
@login_required
def legend():
    return render_template('school_plan/legend.html', active_view='legend', **_form_context())


def _active_plan_classes():
    current_year = AcademicYear.query.filter_by(is_current=True).first()
    query = SchoolClass.query.filter(
        SchoolClass.is_active.is_(True),
        SchoolClass.is_archived.is_(False),
    )
    if current_year:
        query = query.filter(SchoolClass.academic_year_id == current_year.id)
    return query.order_by(SchoolClass.grade.asc().nullslast(), SchoolClass.name.asc()).all()


def _form_context(event=None):
    ensure_school_plan_seed_data()
    classes = _active_plan_classes()
    departments = Department.query.order_by(Department.name.asc()).all()
    class_buildings = sorted(
        {row.building for row in classes if row.building},
        key=lambda row: ((row.short_name or row.name or '').lower(), row.id),
    )
    building_colors = {
        row.id: BUILDING_COLORS[index % len(BUILDING_COLORS)]
        for index, row in enumerate(class_buildings)
    }
    grouped_classes = []
    for grade in sorted({row.grade for row in classes if row.grade is not None}):
        grouped_classes.append(
            {
                'grade': grade,
                'label': f'{grade}-е классы',
                'classes': [row for row in classes if row.grade == grade],
            }
        )
    without_grade = [row for row in classes if row.grade is None]
    if without_grade:
        grouped_classes.append(
            {'grade': None, 'label': 'Без указанной параллели', 'classes': without_grade}
        )

    if request.method == 'POST':
        selected_responsible_ids = {
            int(value) for value in request.form.getlist('responsible_user_ids') if value.isdigit()
        }
        selected_grades = {
            int(value) for value in request.form.getlist('target_grades') if value.isdigit()
        }
        selected_class_ids = {
            int(value) for value in request.form.getlist('target_class_ids') if value.isdigit()
        }
        selected_audience_groups = set(request.form.getlist('audience_groups'))
        selected_responsible_groups = set(request.form.getlist('responsible_groups'))
        audience_scope = request.form.get('audience_scope') or 'school'
        student_target_mode = request.form.get('student_target_mode') or 'all'
        date_mode = request.form.get('date_mode') or 'single'
    else:
        selected_responsible_ids = {
            link.user_id for link in (event.responsible_links if event else [])
        }
        if event and not selected_responsible_ids and event.responsible_user_id:
            selected_responsible_ids.add(event.responsible_user_id)
        selected_grades = {link.grade for link in (event.target_grade_links if event else [])}
        selected_class_ids = {
            link.class_id for link in (event.target_class_links if event else [])
        }
        if event and not selected_grades and not selected_class_ids and event.class_id:
            selected_class_ids.add(event.class_id)
        selected_audience_groups = {
            link.form_value
            for link in (event.group_links if event else [])
            if link.purpose == 'audience'
        }
        selected_responsible_groups = {
            link.form_value
            for link in (event.group_links if event else [])
            if link.purpose == 'responsible'
        }
        if selected_grades or selected_class_ids or 'special:all_students' in selected_audience_groups:
            audience_scope = 'students'
        elif selected_audience_groups:
            audience_scope = 'staff'
        else:
            audience_scope = 'school'
        student_target_mode = (
            'selected' if selected_grades or selected_class_ids else 'all'
        )
        date_mode = (
            'period'
            if event and event.end_date and event.end_date != event.start_date
            else 'single'
        )

    filter_year, filter_month = _selected_month()
    return {
        'directions': SchoolPlanDirection.query.filter_by(is_active=True).order_by(
            SchoolPlanDirection.sort_order.asc(), SchoolPlanDirection.name.asc()
        ).all(),
        'users': User.query.filter(User.is_active_user.is_(True)).order_by(
            User.last_name.asc(), User.first_name.asc(), User.middle_name.asc()
        ).all(),
        'position_options': SCHOOL_PLAN_POSITION_OPTIONS,
        'departments': departments,
        'class_groups': grouped_classes,
        'class_buildings': class_buildings,
        'building_colors': building_colors,
        'selected_responsible_ids': selected_responsible_ids,
        'selected_responsible_groups': selected_responsible_groups,
        'selected_audience_groups': selected_audience_groups,
        'selected_grades': selected_grades,
        'selected_class_ids': selected_class_ids,
        'audience_scope': audience_scope,
        'student_target_mode': student_target_mode,
        'date_mode': date_mode,
        'status_titles': STATUS_TITLES,
        'priority_titles': PRIORITY_TITLES,
        'visibility_titles': VISIBILITY_TITLES,
        'can_manage_editors': _can_manage_editors(),
        'current_args': _request_args_dict(),
        'month_options': _month_filter_options(filter_year, filter_month),
    }


def _int_values(name):
    values = []
    for raw in request.form.getlist(name):
        try:
            value = int(raw)
        except (TypeError, ValueError):
            continue
        if value not in values:
            values.append(value)
    return values


def _group_links_from_form(name, purpose, allowed_specials=()):
    position_labels = dict(SCHOOL_PLAN_POSITION_OPTIONS)
    special_labels = {
        'all_students': 'Все обучающиеся',
        'all_staff': 'Все сотрудники',
        'class_teachers': 'Классные руководители',
    }
    department_ids = []
    parsed = []
    seen = set()
    for raw in request.form.getlist(name):
        group_type, separator, group_key = (raw or '').partition(':')
        if not separator:
            continue
        identity = (group_type, group_key)
        if identity in seen:
            continue
        if group_type == 'special' and group_key in allowed_specials:
            parsed.append((group_type, group_key, special_labels[group_key]))
            seen.add(identity)
        elif group_type == 'role' and group_key in position_labels:
            parsed.append((group_type, group_key, position_labels[group_key]))
            seen.add(identity)
        elif group_type == 'department' and group_key.isdigit():
            department_ids.append(int(group_key))
            seen.add(identity)

    departments = {}
    if department_ids:
        departments = {
            row.id: row
            for row in Department.query.filter(Department.id.in_(department_ids)).all()
        }
        if len(departments) != len(set(department_ids)):
            raise SchoolPlanFormError('Одна из выбранных кафедр недоступна.')
    for department_id in department_ids:
        department = departments[department_id]
        parsed.append(
            ('department', str(department_id), f'Кафедра «{department.name}»')
        )

    return [
        SchoolPlanEventGroup(
            purpose=purpose,
            group_type=group_type,
            group_key=group_key,
            label=label,
            sort_order=index,
        )
        for index, (group_type, group_key, label) in enumerate(parsed)
    ]


def _sync_event_links(collection, desired_links, identity_fields, update_fields=()):
    """Synchronize event child rows without recreating unchanged unique keys."""
    existing_by_identity = {
        tuple(getattr(link, field) for field in identity_fields): link
        for link in collection
    }
    synchronized = []
    for desired_link in desired_links:
        identity = tuple(
            getattr(desired_link, field) for field in identity_fields
        )
        current_link = existing_by_identity.pop(identity, None)
        if current_link is None:
            current_link = desired_link
        else:
            for field in update_fields:
                setattr(current_link, field, getattr(desired_link, field))
        synchronized.append(current_link)
    collection[:] = synchronized


def _apply_event_form(event):
    title = (request.form.get('title') or '').strip()
    if not title:
        raise SchoolPlanFormError('Укажите название мероприятия.')
    if len(title) > 255:
        raise SchoolPlanFormError('Название мероприятия не должно превышать 255 символов.')

    start_date = _parse_date(request.form.get('start_date'))
    if not start_date:
        raise SchoolPlanFormError('Укажите дату мероприятия.')
    start_time = _parse_time(request.form.get('start_time'))
    date_mode = (request.form.get('date_mode') or 'single').strip()
    if date_mode not in {'single', 'period'}:
        raise SchoolPlanFormError('Выберите один день или период.')
    end_date = None
    if date_mode == 'period':
        end_date = _parse_date(request.form.get('end_date'))
        if not end_date:
            raise SchoolPlanFormError('Для периода укажите дату окончания.')
        if end_date < start_date:
            raise SchoolPlanFormError('Дата окончания не может быть раньше даты начала.')

    direction_value = (request.form.get('direction_id') or '').strip()
    direction = None
    new_direction_name = None
    if direction_value == 'other':
        new_direction_name = (request.form.get('direction_other') or '').strip()
        if not new_direction_name:
            raise SchoolPlanFormError('Укажите иное направление деятельности.')
        if len(new_direction_name) > 120:
            raise SchoolPlanFormError('Название направления не должно превышать 120 символов.')
        direction = SchoolPlanDirection.query.filter(
            func.lower(SchoolPlanDirection.name) == new_direction_name.lower()
        ).first()
    elif direction_value.isdigit():
        direction = SchoolPlanDirection.query.filter_by(
            id=int(direction_value), is_active=True
        ).first()
    if not direction and not new_direction_name:
        raise SchoolPlanFormError('Выберите направление деятельности.')

    responsible_ids = _int_values('responsible_user_ids')
    responsible_other = (request.form.get('responsible_other') or '').strip()
    if len(responsible_other) > 255:
        raise SchoolPlanFormError('Поле других ответственных не должно превышать 255 символов.')
    responsible_users = []
    if responsible_ids:
        responsible_users = User.query.filter(
            User.id.in_(responsible_ids), User.is_active_user.is_(True)
        ).all()
        users_by_id = {row.id: row for row in responsible_users}
        if len(users_by_id) != len(responsible_ids):
            raise SchoolPlanFormError('Один из выбранных ответственных недоступен.')
        responsible_users = [users_by_id[user_id] for user_id in responsible_ids]
    responsible_group_links = _group_links_from_form(
        'responsible_groups',
        'responsible',
        allowed_specials={'class_teachers'},
    )

    audience_scope = (request.form.get('audience_scope') or 'school').strip()
    if audience_scope not in {'school', 'students', 'staff'}:
        raise SchoolPlanFormError('Выберите аудиторию мероприятия.')
    selected_grades = []
    selected_classes = []
    audience_group_links = []
    if audience_scope == 'students':
        student_target_mode = (request.form.get('student_target_mode') or 'all').strip()
        if student_target_mode == 'all':
            audience_group_links = [
                SchoolPlanEventGroup(
                    purpose='audience',
                    group_type='special',
                    group_key='all_students',
                    label='Все обучающиеся',
                    sort_order=0,
                )
            ]
        elif student_target_mode == 'selected':
            selected_grades = [grade for grade in _int_values('target_grades') if 1 <= grade <= 11]
            class_ids = _int_values('target_class_ids')
            if class_ids:
                rows = SchoolClass.query.filter(
                    SchoolClass.id.in_(class_ids),
                    SchoolClass.is_active.is_(True),
                    SchoolClass.is_archived.is_(False),
                ).all()
                rows_by_id = {row.id: row for row in rows}
                if len(rows_by_id) != len(class_ids):
                    raise SchoolPlanFormError('Один из выбранных классов недоступен.')
                selected_classes = [rows_by_id[class_id] for class_id in class_ids]
            selected_classes = [
                row for row in selected_classes if row.grade not in selected_grades
            ]
            if not selected_grades and not selected_classes:
                raise SchoolPlanFormError('Выберите хотя бы одну параллель или класс.')
        else:
            raise SchoolPlanFormError('Выберите всех обучающихся или отдельные классы.')
    elif audience_scope == 'staff':
        audience_group_links = _group_links_from_form(
            'audience_groups',
            'audience',
            allowed_specials={'all_staff', 'class_teachers'},
        )
        if any(link.form_value == 'special:all_staff' for link in audience_group_links):
            audience_group_links = [
                link for link in audience_group_links
                if link.form_value == 'special:all_staff'
            ]
        if not audience_group_links:
            raise SchoolPlanFormError('Выберите сотрудников, должности или кафедры.')

    description = (request.form.get('description') or '').strip() or None
    if new_direction_name and not direction:
        direction = SchoolPlanDirection(
            name=new_direction_name,
            color='#475569',
            text_color='#ffffff',
            sort_order=500,
            is_active=True,
        )
        db.session.add(direction)

    event.title = title
    event.description = description
    event.start_date = start_date
    event.start_time = start_time
    event.end_date = end_date
    event.period_type = 'range' if end_date else 'day'
    event.direction = direction
    responsible_links = [
        SchoolPlanEventResponsible(user_id=user.id, sort_order=index)
        for index, user in enumerate(responsible_users)
    ]
    _sync_event_links(
        event.responsible_links,
        responsible_links,
        identity_fields=('user_id',),
        update_fields=('sort_order',),
    )
    event.responsible_user_id = responsible_users[0].id if responsible_users else None
    event.responsible_text = responsible_other or None
    _sync_event_links(
        event.group_links,
        responsible_group_links + audience_group_links,
        identity_fields=('purpose', 'group_type', 'group_key'),
        update_fields=('label', 'sort_order'),
    )
    target_grade_links = [
        SchoolPlanEventGrade(grade=grade) for grade in sorted(selected_grades)
    ]
    _sync_event_links(
        event.target_grade_links,
        target_grade_links,
        identity_fields=('grade',),
    )
    target_class_links = [
        SchoolPlanEventClass(class_id=row.id, sort_order=index)
        for index, row in enumerate(selected_classes)
    ]
    _sync_event_links(
        event.target_class_links,
        target_class_links,
        identity_fields=('class_id',),
        update_fields=('sort_order',),
    )

    if audience_scope in {'school', 'staff'}:
        event.visibility_level = 'school'
        event.class_id = None
        event.building_id = None
    elif audience_scope == 'students':
        event.visibility_level = 'class' if selected_grades or selected_classes else 'school'
        all_target_classes = list(selected_classes)
        if selected_grades:
            all_target_classes.extend(
                SchoolClass.query.filter(
                    SchoolClass.grade.in_(selected_grades),
                    SchoolClass.is_active.is_(True),
                    SchoolClass.is_archived.is_(False),
                ).all()
            )
        event.class_id = all_target_classes[0].id if all_target_classes else None
        building_ids = {
            row.building_id for row in all_target_classes if row.building_id is not None
        }
        event.building_id = next(iter(building_ids)) if len(building_ids) == 1 else None

    current_year = AcademicYear.query.filter_by(is_current=True).first()
    if current_year:
        event.academic_year_id = current_year.id
    if not event.id:
        event.status = 'planned'
        event.priority = 'normal'
        event.created_by_user_id = current_user.id
    event.updated_by_user_id = current_user.id
