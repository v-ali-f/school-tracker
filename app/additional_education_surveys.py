from __future__ import annotations

import io
import os
import re
import secrets
import shutil
import uuid
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace

from flask import (
    Blueprint,
    abort,
    current_app,
    flash,
    g,
    redirect,
    render_template,
    request,
    send_file,
    url_for,
)
from flask_login import current_user, login_required
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.orm import selectinload

from app.core.extensions import db
from app.models import (
    AdditionalEducationSurvey,
    AdditionalEducationSurveyAccess,
    AdditionalEducationSurveyAnswer,
    AdditionalEducationSurveyAttachment,
    AdditionalEducationSurveyQuestion,
    AdditionalEducationSurveySubmission,
    User,
)
from app.models.additional_education_survey import SURVEY_FIELD_TYPES
from app.models.role_access import (
    ACCESS_LEVEL_EDIT,
    ACCESS_LEVEL_FULL,
    ACCESS_LEVEL_HIDDEN,
    ACCESS_LEVEL_VIEW,
    RoleModuleAccess,
)
from app.permissions import _user_role_codes


additional_education_surveys_bp = Blueprint("additional_education_surveys", __name__)

COLLECTION_FORMS_MODULE = "collection_forms"
ACCESS_LEVEL_RANK = {
    ACCESS_LEVEL_HIDDEN: 0,
    ACCESS_LEVEL_VIEW: 1,
    ACCESS_LEVEL_EDIT: 2,
    ACCESS_LEVEL_FULL: 3,
}
ALLOWED_UPLOADS = {"pdf", "doc", "docx", "xls", "xlsx", "odt", "jpg", "jpeg", "png"}
MAX_FILE_SIZE = 10 * 1024 * 1024
MAX_FILES_PER_QUESTION = 5
MAX_TOTAL_UPLOAD_SIZE = 25 * 1024 * 1024
MAX_QUESTIONS = 50
SHORT_CODE_PATTERN = re.compile(r"^[a-z0-9](?:[a-z0-9-]{1,38}[a-z0-9])?$")

STARTER_QUESTIONS = (
    ("ФИО педагога", "short_text", True, (), ""),
    ("Название объединения дополнительного образования", "short_text", True, (), ""),
    ("Максимальное количество детей в группе", "number", True, (), "Укажите целое число."),
    ("Возраст детей", "short_text", True, (), "Например: 7–10 лет."),
    ("Краткое описание (уникальность программы)", "long_text", True, (), ""),
    ("Адрес проведения занятий", "short_text", True, (), ""),
    ("Расписание (день недели, время)", "long_text", True, (), ""),
    ("Кабинет/зал", "short_text", True, (), ""),
    (
        "Объединение откроется после набора группы. Предполагаемая дата начала занятий",
        "date",
        True,
        (),
        "",
    ),
    (
        "Объединение будет работать до конца мая или июня/иной срок?",
        "choice_other",
        True,
        ("До конца мая", "До конца июня"),
        "Если выбираете иной срок, укажите его.",
    ),
    (
        "Документы по программе (при наличии)",
        "file",
        False,
        (),
        "Можно приложить до 5 файлов: PDF, Word, Excel, ODT, JPG или PNG; до 10 МБ каждый.",
    ),
)


def _explicit_role_codes(user):
    codes = {
        str(getattr(role, "code", "")).upper()
        for role in (getattr(user, "roles", None) or [])
        if getattr(role, "code", None)
    }
    primary = getattr(user, "role", None)
    if primary:
        codes.add(str(primary).upper())
    return codes


def collection_forms_access_level(user=None):
    user = user or current_user
    if not user or not getattr(user, "is_authenticated", False):
        return ACCESS_LEVEL_HIDDEN

    explicit_codes = _explicit_role_codes(user)
    role_codes = set(_user_role_codes(user=user)) | explicit_codes
    try:
        preview_codes = getattr(g, "_hub_preview_role_codes", None)
    except RuntimeError:
        preview_codes = None
    if preview_codes is not None:
        explicit_codes = {str(code).upper() for code in preview_codes}
        role_codes = set(explicit_codes)
    try:
        rows = RoleModuleAccess.query.filter(
            RoleModuleAccess.role_code.in_(role_codes),
            RoleModuleAccess.module_code == COLLECTION_FORMS_MODULE,
            RoleModuleAccess.is_active.is_(True),
        ).all()
        rows_by_role = {row.role_code: row for row in rows}
    except Exception:
        rows_by_role = {}

    best_level = ACCESS_LEVEL_HIDDEN
    for role_code in role_codes:
        row = rows_by_role.get(role_code)
        if row is not None:
            level = (
                row.access_level
                if row.is_visible and row.is_enabled
                else ACCESS_LEVEL_HIDDEN
            )
        elif role_code in explicit_codes and role_code in {"ADMIN", "DIRECTOR", "METHODIST"}:
            level = ACCESS_LEVEL_FULL
        else:
            level = ACCESS_LEVEL_HIDDEN
        if ACCESS_LEVEL_RANK.get(level, 0) > ACCESS_LEVEL_RANK.get(best_level, 0):
            best_level = level
    return best_level


def can_access_collection_forms(user=None):
    return ACCESS_LEVEL_RANK.get(collection_forms_access_level(user), 0) >= 1


def _can_create_forms(user=None):
    return ACCESS_LEVEL_RANK.get(collection_forms_access_level(user), 0) >= 2


def _require_section_access(minimum=ACCESS_LEVEL_VIEW):
    if ACCESS_LEVEL_RANK.get(collection_forms_access_level(), 0) < ACCESS_LEVEL_RANK[minimum]:
        abort(403)


def _is_owner(survey, user=None):
    user = user or current_user
    return bool(getattr(user, "id", None) and survey.created_by_user_id == user.id)


def _can_view_results(survey, user=None):
    user = user or current_user
    if _is_owner(survey, user):
        return True
    return bool(
        getattr(user, "id", None)
        and AdditionalEducationSurveyAccess.query.filter_by(
            survey_id=survey.id,
            user_id=user.id,
        ).first()
    )


def _owned_survey_or_404(survey_id):
    return AdditionalEducationSurvey.query.filter_by(
        id=survey_id,
        created_by_user_id=current_user.id,
    ).first_or_404()


def _accessible_survey_or_404(survey_id):
    survey = AdditionalEducationSurvey.query.get_or_404(survey_id)
    if not _can_view_results(survey):
        abort(404)
    return survey


def _public_url(survey):
    if survey.short_code:
        path = url_for("additional_education_surveys.short_public_form", short_code=survey.short_code)
    else:
        path = url_for("additional_education_surveys.public_form", token=survey.public_token)
    base = (current_app.config.get("APP_BASE_URL") or request.url_root).rstrip("/")
    return f"{base}{path}"


def _new_token(model, column):
    while True:
        value = secrets.token_urlsafe(32)
        if not model.query.filter(column == value).first():
            return value


def _parse_short_code(survey=None):
    value = (request.form.get("short_code") or "").strip().lower()
    if not value:
        return None, None
    if not SHORT_CODE_PATTERN.fullmatch(value):
        return None, "Короткое имя должно содержать 3–40 латинских букв, цифр или дефисов."
    query = AdditionalEducationSurvey.query.filter_by(short_code=value)
    if survey is not None and survey.id is not None:
        query = query.filter(AdditionalEducationSurvey.id != survey.id)
    if query.first():
        return None, "Такое короткое имя уже используется другой формой."
    return value, None


def _parse_questions():
    labels = request.form.getlist("question_label[]")
    field_types = request.form.getlist("question_type[]")
    help_texts = request.form.getlist("question_help[]")
    option_rows = request.form.getlist("question_options[]")
    required_indexes = set(request.form.getlist("question_required"))

    if not labels or len(labels) > MAX_QUESTIONS:
        return None, f"Добавьте от 1 до {MAX_QUESTIONS} вопросов."
    if not (len(labels) == len(field_types) == len(help_texts) == len(option_rows)):
        return None, "Не удалось прочитать список вопросов. Обновите страницу и повторите."

    result = []
    for index, raw_label in enumerate(labels):
        label = (raw_label or "").strip()
        field_type = (field_types[index] or "").strip()
        help_text = (help_texts[index] or "").strip()
        if not label:
            return None, f"Укажите текст вопроса № {index + 1}."
        if len(label) > 500 or len(help_text) > 500:
            return None, f"Вопрос № {index + 1} или его подсказка слишком длинные."
        if field_type not in SURVEY_FIELD_TYPES:
            return None, f"У вопроса № {index + 1} выбран неизвестный тип."

        options = []
        if field_type in {"choice", "multiple_choice", "choice_other"}:
            options = [part.strip() for part in option_rows[index].splitlines() if part.strip()]
            options = list(dict.fromkeys(options))
            if not options:
                return None, f"Добавьте варианты ответа для вопроса № {index + 1}."
            if len(options) > 30 or any(len(option) > 200 for option in options):
                return None, f"Слишком много или слишком длинные варианты у вопроса № {index + 1}."

        result.append(
            {
                "label": label,
                "field_type": field_type,
                "help_text": help_text or None,
                "options": options,
                "is_required": str(index) in required_indexes,
            }
        )
    return result, None


def _apply_questions(survey, question_rows):
    survey.questions.clear()
    for index, row in enumerate(question_rows):
        question = AdditionalEducationSurveyQuestion(
            label=row["label"],
            field_type=row["field_type"],
            help_text=row["help_text"],
            is_required=row["is_required"],
            sort_order=index,
        )
        question.set_options(row["options"])
        survey.questions.append(question)


def _posted_editor_questions():
    labels = request.form.getlist("question_label[]")
    field_types = request.form.getlist("question_type[]")
    help_texts = request.form.getlist("question_help[]")
    option_rows = request.form.getlist("question_options[]")
    required_indexes = set(request.form.getlist("question_required"))
    rows = []
    for index, label in enumerate(labels):
        rows.append(
            SimpleNamespace(
                label=label,
                field_type=field_types[index] if index < len(field_types) else "short_text",
                help_text=help_texts[index] if index < len(help_texts) else "",
                options=(
                    [part.strip() for part in option_rows[index].splitlines() if part.strip()]
                    if index < len(option_rows)
                    else []
                ),
                is_required=str(index) in required_indexes,
            )
        )
    return rows


def _file_size(file_storage):
    position = file_storage.stream.tell()
    file_storage.stream.seek(0, os.SEEK_END)
    size = file_storage.stream.tell()
    file_storage.stream.seek(position)
    return size


def _file_signature_is_valid(file_storage, extension):
    position = file_storage.stream.tell()
    file_storage.stream.seek(0)
    header = file_storage.stream.read(8)
    file_storage.stream.seek(position)
    if extension == "pdf":
        return header.startswith(b"%PDF-")
    if extension in {"jpg", "jpeg"}:
        return header.startswith(b"\xff\xd8\xff")
    if extension == "png":
        return header.startswith(b"\x89PNG\r\n\x1a\n")
    if extension in {"docx", "xlsx", "odt"}:
        return header.startswith(b"PK")
    if extension in {"doc", "xls"}:
        return header.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1")
    return False


def _validate_public_answers(survey):
    values = {}
    uploads = {}
    errors = {}
    total_upload_size = 0

    for question in survey.questions:
        key = f"answer_{question.id}"
        if question.field_type == "file":
            files = [item for item in request.files.getlist(key) if item and item.filename]
            uploads[question.id] = []
            if question.is_required and not files:
                errors[question.id] = "Приложите хотя бы один документ."
                continue
            if len(files) > MAX_FILES_PER_QUESTION:
                errors[question.id] = f"Можно приложить не более {MAX_FILES_PER_QUESTION} файлов."
                continue
            for item in files:
                original = Path(item.filename.replace("\\", "/")).name.strip() or "document"
                original = "".join(character for character in original if 31 < ord(character) != 127)
                original = original or "document"
                extension = Path(original).suffix.lower().lstrip(".")
                size = _file_size(item)
                total_upload_size += size
                if extension not in ALLOWED_UPLOADS:
                    errors[question.id] = f"Недопустимый тип файла: {original}."
                    break
                if size <= 0 or size > MAX_FILE_SIZE:
                    errors[question.id] = f"Файл {original} должен быть не больше 10 МБ."
                    break
                if not _file_signature_is_valid(item, extension):
                    errors[question.id] = f"Содержимое файла {original} не соответствует его типу."
                    break
                uploads[question.id].append((item, original[:255], extension, size))
            continue

        if question.field_type == "multiple_choice":
            selected = list(
                dict.fromkeys(
                    item.strip()
                    for item in request.form.getlist(key)
                    if item and item.strip()
                )
            )
            if any(item not in question.options for item in selected):
                errors[question.id] = "Выберите варианты только из предложенного списка."
            elif question.is_required and not selected:
                errors[question.id] = "Выберите хотя бы один вариант."
            value = "\n".join(selected)
            if len(value) > 10000:
                errors[question.id] = "Ответ слишком длинный."
            values[question.id] = value
            continue

        value = (request.form.get(key) or "").strip()
        if question.field_type == "choice_other" and value == "__other__":
            other = (request.form.get(f"answer_other_{question.id}") or "").strip()
            if other:
                value = f"Иной вариант: {other}"
            else:
                errors[question.id] = "Укажите иной вариант."
        elif question.field_type in {"choice", "choice_other"} and value:
            if value not in question.options:
                errors[question.id] = "Выберите один из предложенных вариантов."
        elif question.field_type == "yes_no" and value:
            if value not in {"Да", "Нет"}:
                errors[question.id] = "Выберите «Да» или «Нет»."
        elif question.field_type == "number" and value:
            try:
                number = int(value)
                if number < 0:
                    raise ValueError
            except ValueError:
                errors[question.id] = "Укажите целое неотрицательное число."
        elif question.field_type == "date" and value:
            try:
                datetime.strptime(value, "%Y-%m-%d")
            except ValueError:
                errors[question.id] = "Укажите корректную дату."
        if question.is_required and not value:
            errors[question.id] = "Это обязательный вопрос."
        if len(value) > 10000:
            errors[question.id] = "Ответ слишком длинный."
        values[question.id] = value

    if total_upload_size > MAX_TOTAL_UPLOAD_SIZE:
        errors["files"] = "Общий размер документов не должен превышать 25 МБ."
    return values, uploads, errors


def _save_submission(survey, values, uploads):
    submission = AdditionalEducationSurveySubmission(
        survey_id=survey.id,
        receipt_token=_new_token(
            AdditionalEducationSurveySubmission,
            AdditionalEducationSurveySubmission.receipt_token,
        ),
    )
    db.session.add(submission)
    db.session.flush()

    saved_paths = []
    try:
        for question in survey.questions:
            answer = AdditionalEducationSurveyAnswer(
                submission_id=submission.id,
                question_id=question.id,
                value_text=values.get(question.id) or None,
            )
            db.session.add(answer)
            db.session.flush()
            for file_storage, original, extension, size in uploads.get(question.id, []):
                relative = Path("additional_education_surveys") / str(survey.id) / submission.receipt_token
                destination_dir = Path(current_app.config["UPLOAD_FOLDER"]) / relative
                destination_dir.mkdir(parents=True, exist_ok=True)
                stored_name = f"{uuid.uuid4().hex}.{extension}"
                destination = destination_dir / stored_name
                file_storage.stream.seek(0)
                file_storage.save(destination)
                saved_paths.append(destination)
                db.session.add(
                    AdditionalEducationSurveyAttachment(
                        answer_id=answer.id,
                        original_filename=original,
                        stored_path=str(relative / stored_name),
                        content_type=file_storage.mimetype,
                        file_size=size,
                    )
                )
        db.session.commit()
        return submission
    except Exception:
        db.session.rollback()
        for path in saved_paths:
            try:
                path.unlink(missing_ok=True)
            except OSError:
                pass
        raise


@additional_education_surveys_bp.route("/additional-education/surveys/")
@additional_education_surveys_bp.route("/collection-forms/")
@login_required
def index():
    _require_section_access()
    owned_surveys = (
        AdditionalEducationSurvey.query
        .filter_by(created_by_user_id=current_user.id)
        .order_by(AdditionalEducationSurvey.created_at.desc())
        .all()
    )
    shared_surveys = (
        AdditionalEducationSurvey.query
        .join(AdditionalEducationSurveyAccess)
        .filter(AdditionalEducationSurveyAccess.user_id == current_user.id)
        .order_by(AdditionalEducationSurvey.created_at.desc())
        .all()
    )

    def row_for(survey, is_owner):
        return {
            "survey": survey,
            "response_count": survey.submissions.count(),
            "public_url": _public_url(survey),
            "is_owner": is_owner,
        }

    return render_template(
        "additional_education_surveys/index.html",
        owned_rows=[row_for(survey, True) for survey in owned_surveys],
        shared_rows=[row_for(survey, False) for survey in shared_surveys],
        can_create_forms=_can_create_forms(),
    )


@additional_education_surveys_bp.route("/additional-education/surveys/new", methods=["GET", "POST"])
@additional_education_surveys_bp.route("/collection-forms/new", methods=["GET", "POST"])
@login_required
def new():
    _require_section_access(ACCESS_LEVEL_EDIT)
    if request.method == "POST":
        title = (request.form.get("title") or "").strip()
        question_rows, error = _parse_questions()
        short_code, short_code_error = _parse_short_code()
        if not title:
            error = "Укажите название формы."
        if len(title) > 255:
            error = "Название формы слишком длинное."
        if short_code_error:
            error = short_code_error
        if error:
            flash(error, "danger")
        else:
            survey = AdditionalEducationSurvey(
                title=title,
                description=(request.form.get("description") or "").strip() or None,
                public_token=_new_token(AdditionalEducationSurvey, AdditionalEducationSurvey.public_token),
                short_code=short_code,
                is_published=request.form.get("is_published") == "1",
                created_by_user_id=current_user.id,
            )
            _apply_questions(survey, question_rows)
            db.session.add(survey)
            db.session.commit()
            flash("Форма создана.", "success")
            return redirect(url_for("additional_education_surveys.index"))
    return render_template(
        "additional_education_surveys/editor.html",
        survey=None,
        field_types=SURVEY_FIELD_TYPES,
        locked=False,
        editor_questions=_posted_editor_questions() if request.method == "POST" else None,
    )


@additional_education_surveys_bp.post("/additional-education/surveys/from-template")
@additional_education_surveys_bp.post("/collection-forms/from-template")
@login_required
def create_from_template():
    _require_section_access(ACCESS_LEVEL_EDIT)
    survey = AdditionalEducationSurvey(
        title="Открытие записи на mos.ru (внебюджет)",
        description="Заполните сведения об объединении дополнительного образования.",
        public_token=_new_token(AdditionalEducationSurvey, AdditionalEducationSurvey.public_token),
        is_published=True,
        created_by_user_id=current_user.id,
    )
    rows = [
        {
            "label": label,
            "field_type": field_type,
            "is_required": required,
            "options": options,
            "help_text": help_text or None,
        }
        for label, field_type, required, options, help_text in STARTER_QUESTIONS
    ]
    _apply_questions(survey, rows)
    db.session.add(survey)
    db.session.commit()
    flash("Готовая форма создана и открыта для ответов.", "success")
    return redirect(url_for("additional_education_surveys.edit", survey_id=survey.id))


@additional_education_surveys_bp.route(
    "/additional-education/surveys/<int:survey_id>/edit", methods=["GET", "POST"]
)
@additional_education_surveys_bp.route(
    "/collection-forms/<int:survey_id>/edit", methods=["GET", "POST"]
)
@login_required
def edit(survey_id):
    _require_section_access(ACCESS_LEVEL_EDIT)
    survey = _owned_survey_or_404(survey_id)
    locked = survey.submissions.count() > 0
    if request.method == "POST":
        title = (request.form.get("title") or "").strip()
        short_code, short_code_error = _parse_short_code(survey)
        if not title or len(title) > 255:
            flash("Укажите название формы длиной до 255 символов.", "danger")
        elif short_code_error:
            flash(short_code_error, "danger")
        else:
            if not locked:
                question_rows, error = _parse_questions()
                if error:
                    flash(error, "danger")
                    return render_template(
                        "additional_education_surveys/editor.html",
                        survey=survey,
                        field_types=SURVEY_FIELD_TYPES,
                        locked=locked,
                        public_url=_public_url(survey),
                    )
                _apply_questions(survey, question_rows)
            survey.title = title
            survey.description = (request.form.get("description") or "").strip() or None
            survey.short_code = short_code
            survey.is_published = request.form.get("is_published") == "1"
            db.session.commit()
            flash("Форма сохранена.", "success")
            return redirect(url_for("additional_education_surveys.edit", survey_id=survey.id))
    return render_template(
        "additional_education_surveys/editor.html",
        survey=survey,
        field_types=SURVEY_FIELD_TYPES,
        locked=locked,
        public_url=_public_url(survey),
    )


@additional_education_surveys_bp.post("/additional-education/surveys/<int:survey_id>/toggle")
@additional_education_surveys_bp.post("/collection-forms/<int:survey_id>/toggle")
@login_required
def toggle(survey_id):
    _require_section_access(ACCESS_LEVEL_EDIT)
    survey = _owned_survey_or_404(survey_id)
    survey.is_published = not survey.is_published
    db.session.commit()
    flash("Приём ответов открыт." if survey.is_published else "Приём ответов остановлен.", "success")
    return redirect(url_for("additional_education_surveys.index"))


@additional_education_surveys_bp.route("/additional-education/surveys/<int:survey_id>/results")
@additional_education_surveys_bp.route("/collection-forms/<int:survey_id>/results")
@login_required
def results(survey_id):
    _require_section_access()
    survey = _accessible_survey_or_404(survey_id)
    submissions = (
        AdditionalEducationSurveySubmission.query
        .filter_by(survey_id=survey.id)
        .options(selectinload(AdditionalEducationSurveySubmission.answers))
        .order_by(AdditionalEducationSurveySubmission.submitted_at.desc())
        .all()
    )
    return render_template(
        "additional_education_surveys/results.html",
        survey=survey,
        submissions=submissions,
        public_url=_public_url(survey),
        is_owner=_is_owner(survey),
    )


@additional_education_surveys_bp.route(
    "/additional-education/surveys/<int:survey_id>/submissions/<int:submission_id>"
)
@additional_education_surveys_bp.route(
    "/collection-forms/<int:survey_id>/submissions/<int:submission_id>"
)
@login_required
def submission_detail(survey_id, submission_id):
    _require_section_access()
    survey = _accessible_survey_or_404(survey_id)
    submission = AdditionalEducationSurveySubmission.query.filter_by(
        id=submission_id, survey_id=survey.id
    ).first_or_404()
    return render_template(
        "additional_education_surveys/submission_detail.html",
        survey=survey,
        submission=submission,
        answers=submission.answers_by_question_id,
    )


@additional_education_surveys_bp.route(
    "/additional-education/surveys/<int:survey_id>/export.xlsx"
)
@additional_education_surveys_bp.route(
    "/collection-forms/<int:survey_id>/export.xlsx"
)
@login_required
def export_results(survey_id):
    _require_section_access()
    survey = _accessible_survey_or_404(survey_id)
    submissions = (
        AdditionalEducationSurveySubmission.query
        .filter_by(survey_id=survey.id)
        .order_by(AdditionalEducationSurveySubmission.submitted_at.asc())
        .all()
    )
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Ответы"
    headers = ["№", "Дата и время"] + [question.label for question in survey.questions]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2457A7")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row_number, submission in enumerate(submissions, start=1):
        answers = submission.answers_by_question_id
        values = [row_number, submission.submitted_at.strftime("%d.%m.%Y %H:%M")]
        for question in survey.questions:
            answer = answers.get(question.id)
            if not answer:
                values.append("")
            elif question.field_type == "file":
                values.append("; ".join(item.original_filename for item in answer.attachments))
            else:
                values.append(answer.value_text or "")
        sheet.append(values)
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        letter = column[0].column_letter
        content_width = max(len(str(cell.value or "")) for cell in column) + 2
        sheet.column_dimensions[letter].width = min(50, max(12, content_width))
        for cell in column:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(
        output,
        as_attachment=True,
        download_name=f"additional_education_survey_{survey.id}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@additional_education_surveys_bp.route(
    "/additional-education/surveys/attachments/<int:attachment_id>/download"
)
@additional_education_surveys_bp.route(
    "/collection-forms/attachments/<int:attachment_id>/download"
)
@login_required
def download_attachment(attachment_id):
    _require_section_access()
    attachment = AdditionalEducationSurveyAttachment.query.get_or_404(attachment_id)
    survey = attachment.answer.submission.survey
    if not _can_view_results(survey):
        abort(404)
    upload_root = Path(current_app.config["UPLOAD_FOLDER"]).resolve()
    path = (upload_root / attachment.stored_path).resolve()
    if upload_root not in path.parents or not path.is_file():
        abort(404)
    return send_file(
        path,
        as_attachment=True,
        download_name=attachment.original_filename,
        mimetype="application/octet-stream",
    )


@additional_education_surveys_bp.route(
    "/collection-forms/<int:survey_id>/sharing",
    methods=["GET", "POST"],
)
@login_required
def sharing(survey_id):
    _require_section_access(ACCESS_LEVEL_EDIT)
    survey = _owned_survey_or_404(survey_id)
    users = (
        User.query
        .filter(
            User.id != current_user.id,
            User.is_active_user.is_(True),
            User.archived_at.is_(None),
        )
        .order_by(User.last_name.asc(), User.first_name.asc(), User.username.asc())
        .all()
    )
    allowed_user_ids = {user.id for user in users}

    if request.method == "POST":
        selected_user_ids = set()
        for raw_user_id in request.form.getlist("user_ids"):
            try:
                user_id = int(raw_user_id)
            except (TypeError, ValueError):
                continue
            if user_id in allowed_user_ids:
                selected_user_ids.add(user_id)

        existing = {grant.user_id: grant for grant in survey.access_grants}
        for user_id, grant in existing.items():
            if user_id not in selected_user_ids:
                db.session.delete(grant)
        for user_id in selected_user_ids - set(existing):
            db.session.add(
                AdditionalEducationSurveyAccess(
                    survey_id=survey.id,
                    user_id=user_id,
                    granted_by_user_id=current_user.id,
                )
            )
        db.session.commit()
        flash("Доступ к результатам обновлён.", "success")
        return redirect(url_for("additional_education_surveys.sharing", survey_id=survey.id))

    return render_template(
        "additional_education_surveys/sharing.html",
        survey=survey,
        users=users,
        selected_user_ids={grant.user_id for grant in survey.access_grants},
    )


@additional_education_surveys_bp.post("/collection-forms/<int:survey_id>/delete")
@login_required
def delete(survey_id):
    _require_section_access(ACCESS_LEVEL_EDIT)
    survey = _owned_survey_or_404(survey_id)
    upload_root = Path(current_app.config["UPLOAD_FOLDER"]).resolve()
    survey_upload_dir = (upload_root / "additional_education_surveys" / str(survey.id)).resolve()
    title = survey.title
    db.session.delete(survey)
    db.session.commit()

    if upload_root in survey_upload_dir.parents and survey_upload_dir.is_dir():
        try:
            shutil.rmtree(survey_upload_dir)
        except OSError:
            current_app.logger.exception(
                "Не удалось удалить каталог файлов формы %s",
                survey_id,
            )
    flash(f"Форма «{title}» удалена.", "success")
    return redirect(url_for("additional_education_surveys.index"))


def _serve_public_form(survey, thank_you_url):
    if not survey.is_published:
        return render_template("additional_education_surveys/public_closed.html", survey=survey), 410
    errors = {}
    if request.method == "POST":
        if request.content_length and request.content_length > MAX_TOTAL_UPLOAD_SIZE + 1024 * 1024:
            return render_template(
                "additional_education_surveys/public_form.html",
                survey=survey,
                errors={"files": "Общий размер отправки не должен превышать 25 МБ."},
                allowed_extensions=", ".join(sorted(ALLOWED_UPLOADS)),
                max_file_size_mb=MAX_FILE_SIZE // 1024 // 1024,
            ), 413
        # Невидимое поле отсекает примитивных ботов без сбора персональных сетевых данных.
        if request.form.get("website"):
            return render_template("additional_education_surveys/public_thank_you.html", survey=survey)
        values, uploads, errors = _validate_public_answers(survey)
        if not errors:
            _save_submission(survey, values, uploads)
            return redirect(thank_you_url)
    return render_template(
        "additional_education_surveys/public_form.html",
        survey=survey,
        errors=errors,
        allowed_extensions=", ".join(sorted(ALLOWED_UPLOADS)),
        max_file_size_mb=MAX_FILE_SIZE // 1024 // 1024,
    )


@additional_education_surveys_bp.route(
    "/forms/additional-education/<token>", methods=["GET", "POST"]
)
def public_form(token):
    survey = AdditionalEducationSurvey.query.filter_by(public_token=token).first_or_404()
    return _serve_public_form(
        survey,
        url_for("additional_education_surveys.public_thank_you", token=token),
    )


@additional_education_surveys_bp.route("/s/<short_code>", methods=["GET", "POST"])
def short_public_form(short_code):
    survey = AdditionalEducationSurvey.query.filter_by(short_code=short_code.lower()).first_or_404()
    return _serve_public_form(
        survey,
        url_for("additional_education_surveys.short_public_thank_you", short_code=survey.short_code),
    )


@additional_education_surveys_bp.route("/forms/additional-education/<token>/thank-you")
def public_thank_you(token):
    survey = AdditionalEducationSurvey.query.filter_by(public_token=token).first_or_404()
    return render_template("additional_education_surveys/public_thank_you.html", survey=survey)


@additional_education_surveys_bp.route("/s/<short_code>/thank-you")
def short_public_thank_you(short_code):
    survey = AdditionalEducationSurvey.query.filter_by(short_code=short_code.lower()).first_or_404()
    return render_template("additional_education_surveys/public_thank_you.html", survey=survey)
